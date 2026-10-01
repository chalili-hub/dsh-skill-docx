#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
docx_kit —— 中文实验报告排版工具库（基于 python-docx）

处理好了 python-docx 最常见的坑：中文字体必须显式设置 w:rFonts/@w:eastAsia。

用法::

    from docx_kit import Doc
    d = Doc()                      # A4 + 2.6cm 边距 + 宋体正文
    d.cover(title="...", subtitle="...", meta=[("姓　　名", ""), ("学　　号", "...")])
    d.h("一、实验目的", 1)
    d.p("正文段落，自动首行缩进。")
    d.code(open("cmd.txt").read())
    d.table(["地址", "机器码"], [["0x200", "01400313"]], widths=[3, 4])
    d.pic("fig1.png")
    d.caption("图 1  说明文字")
    d.save("报告.docx")
"""

import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# 中文字体方案：标题黑体 / 正文宋体 / 代码等宽
FONT_HEI = "黑体"
FONT_SONG = "宋体"
FONT_KAI = "楷体"
FONT_MONO = "Consolas"
FONT_ASCII = "Times New Roman"
FONT_ASCII_BOLD = "Arial"

H1_SIZE = 15
H2_SIZE = 13
H3_SIZE = 11.5
BODY_SIZE = 10.5


def set_font(run, ea=FONT_SONG, ascii_f=FONT_ASCII):
    """同时设置西文与中文字体。只设 run.font.name 对中文无效。"""
    run.font.name = ascii_f
    rPr = run._element.get_or_add_rPr()
    rF = rPr.find(qn("w:rFonts"))
    if rF is None:
        rF = OxmlElement("w:rFonts")
        rPr.append(rF)
    rF.set(qn("w:ascii"), ascii_f)
    rF.set(qn("w:hAnsi"), ascii_f)
    rF.set(qn("w:eastAsia"), ea)


def shade(paragraph, fill="F5F5F5"):
    pPr = paragraph._element.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    pPr.append(sh)


def cell_shade(cell, fill="DCE6F1"):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def _looks_numeric(s):
    s = str(s)
    if s.startswith(("0x", "0X")):
        return True
    if len(s) == 8 and all(c in "0123456789abcdefABCDEF" for c in s):
        return True
    return False


class Doc:
    def __init__(self, margin_cm=2.6, top_cm=2.4, bottom_cm=2.4):
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(margin_cm)
        sec.right_margin = Cm(margin_cm)
        sec.top_margin = Cm(top_cm)
        sec.bottom_margin = Cm(bottom_cm)
        # 可用正文宽度：插图与表格按此对齐
        self.text_width_cm = 21.0 - 2 * margin_cm

        st = self.doc.styles["Normal"]
        st.font.name = FONT_ASCII
        st.font.size = Pt(BODY_SIZE)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_SONG)
        st.paragraph_format.line_spacing = 1.35
        st.paragraph_format.space_after = Pt(3)

    # ------------------------------------------------------------------ 文本
    def p(self, text="", size=BODY_SIZE, bold=False, ea=FONT_SONG,
          ascii_f=FONT_ASCII, align=None, indent=True, space_after=3,
          color=None, line=1.35):
        para = self.doc.add_paragraph()
        if indent:
            para.paragraph_format.first_line_indent = Pt(size * 2)
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.line_spacing = line
        if align is not None:
            para.alignment = align
        if text:
            r = para.add_run(text)
            r.font.size = Pt(size)
            r.bold = bold
            if color:
                r.font.color.rgb = color
            set_font(r, ea, ascii_f)
        return para

    def h(self, text, level=1):
        sizes = {1: H1_SIZE, 2: H2_SIZE, 3: H3_SIZE}
        para = self.doc.add_paragraph()
        para.paragraph_format.space_before = Pt(12 if level == 1 else 9)
        para.paragraph_format.space_after = Pt(6)
        r = para.add_run(text)
        r.font.size = Pt(sizes.get(level, H3_SIZE))
        r.bold = True
        set_font(r, FONT_HEI, FONT_ASCII_BOLD)
        return para

    def code(self, text, size=8.5, fill="F5F5F5"):
        """逐行插入代码块（一行一段落，底纹才准确）。"""
        lines = text.strip("\n").split("\n")
        for i, ln in enumerate(lines):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.space_before = Pt(4 if i == 0 else 0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.5)
            r = para.add_run(ln if ln else " ")
            r.font.size = Pt(size)
            set_font(r, FONT_SONG, FONT_MONO)
            shade(para, fill)
            if i == len(lines) - 1:
                pf.space_after = Pt(6)

    def caption(self, text):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.space_after = Pt(9)
        r = para.add_run(text)
        r.font.size = Pt(9)
        r.bold = True
        set_font(r, FONT_SONG, FONT_ASCII)
        return para

    def page_break(self):
        self.doc.add_page_break()

    def blank(self, n=1):
        for _ in range(n):
            self.doc.add_paragraph()

    # ------------------------------------------------------------------ 表格
    def table(self, headers, rows, widths=None, size=9, header_fill="DCE6F1"):
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        hdr = t.rows[0].cells
        for i, txt in enumerate(headers):
            hdr[i].text = ""
            p = hdr[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(txt))
            r.font.size = Pt(size)
            r.bold = True
            set_font(r, FONT_HEI, FONT_ASCII_BOLD)
            cell_shade(hdr[i], header_fill)

        for row in rows:
            cells = t.add_row().cells
            for i, txt in enumerate(row):
                cells[i].text = ""
                p = cells[i].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                r = p.add_run(str(txt))
                r.font.size = Pt(size)
                set_font(r, FONT_SONG, FONT_MONO if _looks_numeric(txt) else FONT_ASCII)

        if widths:
            for i, w in enumerate(widths):
                for row in t.rows:
                    row.cells[i].width = Cm(w)
        return t

    # ------------------------------------------------------------------ 图片
    def pic(self, path, width_cm=None):
        width_cm = width_cm or self.text_width_cm
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(6)
        para.paragraph_format.space_after = Pt(0)
        para.add_run().add_picture(path, width=Cm(width_cm))
        return para

    # ------------------------------------------------------------------ 封面
    def cover(self, title, subtitle="", course="", meta=None, gap_top=3, title_size=26):
        self.blank(gap_top)
        if course:
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(course)
            r.font.size = Pt(15)
            r.bold = True
            set_font(r, FONT_HEI, FONT_ASCII_BOLD)

        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_before = Pt(24)
        r = para.add_run(title)
        r.font.size = Pt(title_size)
        r.bold = True
        set_font(r, FONT_HEI, FONT_ASCII_BOLD)

        if subtitle:
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = para.add_run(subtitle)
            r.font.size = Pt(12)
            set_font(r, FONT_KAI, FONT_ASCII)

        self.blank(4)

        if meta:
            t = self.doc.add_table(rows=len(meta), cols=2)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for i, (k, v) in enumerate(meta):
                c = t.rows[i].cells
                c[0].text = ""
                r = c[0].paragraphs[0].add_run(k + "：")
                r.font.size = Pt(12)
                set_font(r, FONT_HEI, FONT_ASCII_BOLD)
                c[0].width = Cm(3.5)
                c[1].text = ""
                r = c[1].paragraphs[0].add_run(v)
                r.font.size = Pt(12)
                set_font(r, FONT_SONG, FONT_ASCII)
                c[1].width = Cm(9)
        self.page_break()

    # ------------------------------------------------------------------ 保存
    def save(self, path):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.doc.save(path)
        return path
